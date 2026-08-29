"""Importa el Excel de La Triple L y busca fotos por código de barras (Open Food/Beauty/Products Facts)."""
from __future__ import annotations

import json
import os
import re
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = Path(r"c:\Users\leord\Downloads\productos filtrados 25-8.xlsx")
CACHE = ROOT / "data" / "off-cache.json"
PRODUCTOS = ROOT / "data" / "productos.json"
RUBROS = ROOT / "data" / "rubros.json"
RUBRO_IMG_DIR = ROOT / "recursos" / "imagenes-rubros"
PUBLIC_RUBRO = ROOT / "public" / "recursos" / "imagenes-rubros"

UA = "LaTripleLMinimarket/1.0 (catalogo; contacto local)"

SECCION_LABEL = {
    "ALMACEN GENERAL": "Almacén",
    "BEBIDAS": "Bebidas",
    "GALLETAS": "Galletas",
    "LIMPIEZA E HIGENE": "Limpieza e higiene",
    "CIGARRILLOS": "Cigarrillos",
    "VAPERS": "Vapers",
    "ELECTRONICA": "Electrónica",
    "TECNOLOGIA": "Tecnología",
    "OTROS": "Otros",
    "VARIOS": "Varios",
}

FALLBACK_URLS = {
    "ALMACEN GENERAL": "https://images.unsplash.com/photo-1583258292688-d0213dc5a3a8?auto=format&fit=crop&w=900&q=80",
    "BEBIDAS": "https://images.unsplash.com/photo-1629203851122-3726ecdf080e?auto=format&fit=crop&w=900&q=80",
    "GALLETAS": "https://images.unsplash.com/photo-1558961363-fa8fdf82db35?auto=format&fit=crop&w=900&q=80",
    "LIMPIEZA E HIGENE": "https://images.unsplash.com/photo-1563453392212-326f5e854473?auto=format&fit=crop&w=900&q=80",
    "CIGARRILLOS": "https://images.unsplash.com/photo-1600880292203-757bb62b4baf?auto=format&fit=crop&w=900&q=80",
    "VAPERS": "https://images.unsplash.com/photo-1518770660439-4636190af475?auto=format&fit=crop&w=900&q=80",
    "ELECTRONICA": "https://images.unsplash.com/photo-1518770660439-4636190af475?auto=format&fit=crop&w=900&q=80",
    "TECNOLOGIA": "https://images.unsplash.com/photo-1511707171634-5f897ff02aa9?auto=format&fit=crop&w=900&q=80",
    "OTROS": "https://images.unsplash.com/photo-1542838132-92c53300491e?auto=format&fit=crop&w=900&q=80",
    "VARIOS": "https://images.unsplash.com/photo-1542838132-92c53300491e?auto=format&fit=crop&w=900&q=80",
}

APIS = [
    "https://world.openfoodfacts.org/api/v2/product/{code}.json?fields=image_front_url,image_url,image_front_small_url,product_name",
    "https://world.openbeautyfacts.org/api/v2/product/{code}.json?fields=image_front_url,image_url,image_front_small_url,product_name",
    "https://world.openproductsfacts.org/api/v2/product/{code}.json?fields=image_front_url,image_url,image_front_small_url,product_name",
]


def clave_rubro(nombre: str) -> str:
    nfd = nombre.encode("utf-8").decode("utf-8")
    nfd = re.sub(r"[\u0300-\u036f]", "", nfd.encode("ascii", "ignore").decode("ascii") if False else nombre)
    # Keep Spanish letters; uppercase like the TS claveRubro (strip accents)
    import unicodedata

    s = unicodedata.normalize("NFD", nombre)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = s.strip().upper()
    s = re.sub(r"[^A-Z0-9]+", " ", s).strip()
    s = re.sub(r"\s+", " ", s)
    return s


def get_json(url: str, timeout: float = 12) -> dict | None:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "application/json"})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            return json.loads(res.read().decode("utf-8", "replace"))
    except Exception:
        return None


def download(url: str, dest: Path, timeout: float = 20) -> bool:
    dest.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            data = res.read()
        if len(data) < 800:
            return False
        dest.write_bytes(data)
        return True
    except Exception:
        return False


def lookup_barcode(code: str, cache: dict) -> str:
    if code in cache:
        return cache[code] or ""
    digits = re.sub(r"\D", "", code)
    if len(digits) < 8:
        cache[code] = ""
        return ""
    url_img = ""
    for tmpl in APIS:
        payload = get_json(tmpl.format(code=digits))
        if not payload or payload.get("status") != 1:
            continue
        prod = payload.get("product") or {}
        url_img = (
            prod.get("image_front_url")
            or prod.get("image_url")
            or prod.get("image_front_small_url")
            or ""
        )
        if url_img:
            break
    cache[code] = url_img
    return url_img


def fallback_path(seccion: str) -> str:
    slug = seccion.lower().replace(" ", "-")
    return f"/recursos/imagenes-rubros/{slug}.jpg"


def main() -> None:
    RUBRO_IMG_DIR.mkdir(parents=True, exist_ok=True)
    PUBLIC_RUBRO.mkdir(parents=True, exist_ok=True)

    print("Descargando fotos de rubro…")
    for clave, url in FALLBACK_URLS.items():
        slug = clave.lower().replace(" ", "-")
        dest = RUBRO_IMG_DIR / f"{slug}.jpg"
        if dest.exists() and dest.stat().st_size > 800:
            pass
        else:
            ok = download(url, dest)
            print(f"  {clave}: {'OK' if ok else 'FALLÓ'}")
        public = PUBLIC_RUBRO / f"{slug}.jpg"
        if dest.exists():
            public.write_bytes(dest.read_bytes())

    cache: dict[str, str] = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    rows = []
    for i, row in enumerate(ws.iter_rows(min_row=11, max_row=ws.max_row, values_only=True), 11):
        codigo, barras, _, nombre, *_rest = row
        precio = row[8]
        seccion = row[9]
        if not nombre or precio is None:
            continue
        try:
            precio_n = float(precio)
        except (TypeError, ValueError):
            continue
        rows.append(
            {
                "row": i,
                "codigo": str(codigo).strip() if codigo is not None else str(i),
                "barras": str(barras).strip() if barras else "",
                "nombre": str(nombre).strip(),
                "precio": round(precio_n),
                "seccion": clave_rubro(str(seccion or "Otros")),
            }
        )

    barcodes = sorted({r["barras"] for r in rows if r["barras"] and r["barras"] not in cache})
    print(f"Productos: {len(rows)}. Barcodes a consultar: {len(barcodes)} (cache {len(cache)})")

    def worker(code: str) -> tuple[str, str]:
        return code, lookup_barcode(code, {})

    done = 0
    with ThreadPoolExecutor(max_workers=8) as pool:
        futs = {pool.submit(worker, c): c for c in barcodes}
        for fut in as_completed(futs):
            code, url = fut.result()
            cache[code] = url
            done += 1
            if done % 40 == 0 or done == len(barcodes):
                print(f"  fotos {done}/{len(barcodes)}")
                CACHE.parent.mkdir(parents=True, exist_ok=True)
                CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    productos = []
    usadas = set()
    for r in rows:
        seccion = r["seccion"] or "OTROS"
        img = cache.get(r["barras"]) or fallback_path(seccion)
        pid = f"p{r['row']}"
        codigo = r["codigo"] or pid
        if codigo in usadas:
            codigo = f"{codigo}-{r['row']}"
        usadas.add(codigo)
        desc = ""
        if r["barras"]:
            desc = f"EAN {r['barras']}"
        productos.append(
            {
                "id": pid,
                "codigo": codigo,
                "nombre": r["nombre"],
                "seccion": seccion,
                "precio": r["precio"],
                "imagen": img,
                "descripcion": desc,
            }
        )

    claves = []
    seen = set()
    for p in productos:
        if p["seccion"] not in seen:
            seen.add(p["seccion"])
            claves.append(p["seccion"])

    rubros = []
    for clave in claves:
        slug = clave.lower().replace(" ", "-")
        rubros.append(
            {
                "clave": clave,
                "etiqueta": SECCION_LABEL.get(clave) or clave.title(),
                "imagen": f"/recursos/imagenes-rubros/{slug}.jpg",
            }
        )

    PRODUCTOS.parent.mkdir(parents=True, exist_ok=True)
    PRODUCTOS.write_text(json.dumps(productos, ensure_ascii=False, indent=2), encoding="utf-8")
    RUBROS.write_text(json.dumps(rubros, ensure_ascii=False, indent=2), encoding="utf-8")
    con_foto = sum(1 for p in productos if p["imagen"].startswith("http"))
    print(f"Listo: {len(productos)} productos, {con_foto} con foto de catálogo abierto, {len(rubros)} rubros.")


if __name__ == "__main__":
    main()
