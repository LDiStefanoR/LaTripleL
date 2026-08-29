"""Completa fotos del catálogo desde supermercados (VTEX: Jumbo, Disco, Vea, Carrefour)."""
from __future__ import annotations

import json
import re
import unicodedata
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
EXCEL = Path(r"c:\Users\leord\Downloads\productos filtrados 25-8.xlsx")
PRODUCTOS = ROOT / "data" / "productos.json"
CACHE = ROOT / "data" / "super-cache.json"
EAN_MAP = ROOT / "data" / "ean-map.json"
DEST = ROOT / "recursos" / "imagenes-productos"
PUBLIC = ROOT / "public" / "recursos" / "imagenes-productos"

UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

STORES = [
    "https://www.jumbo.com.ar",
    "https://www.disco.com.ar",
    "https://www.vea.com.ar",
    "https://www.carrefour.com.ar",
    "https://diaonline.supermercadosdia.com.ar",
    "https://www.farmacity.com",
]

STOP = {
    "x", "gr", "g", "ml", "lt", "l", "kg", "un", "u", "pack", "paq", "lata",
    "doypack", "frasco", "nvo", "nuevo", "c", "de", "con", "sin", "para",
    "sabor", "original", "clasico", "clásico", "the", "and",
}


def get_json(url: str, timeout: float = 14):
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": UA,
            "Accept": "application/json,text/plain,*/*",
            "Accept-Language": "es-AR,es;q=0.9",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=timeout) as res:
            return json.loads(res.read().decode("utf-8", "replace"))
    except Exception:
        return None


def imagen_vtex(payload) -> str:
    if not isinstance(payload, list) or not payload:
        return ""
    items = payload[0].get("items") or []
    if not items:
        return ""
    images = items[0].get("images") or []
    if not images:
        return ""
    url = images[0].get("imageUrl") or images[0].get("imageUrl") or ""
    return str(url).split("?")[0] if url else ""


def variantes_ean(raw: str) -> list[str]:
    digits = re.sub(r"\D", "", raw or "")
    if len(digits) < 8:
        return []
    out = [digits]
    if len(digits) < 13:
        out.append(digits.zfill(13))
    if len(digits) == 12:
        out.append(digits.zfill(13))
    # unique preserve order
    seen = []
    for c in out:
        if c not in seen:
            seen.append(c)
    return seen


def buscar_ean(ean: str) -> str:
    for code in variantes_ean(ean):
        for store in STORES:
            data = get_json(
                f"{store}/api/catalog_system/pub/products/search?fq=alternateIds_Ean:{code}"
            )
            url = imagen_vtex(data)
            if url:
                return url
    return ""


def tokens(texto: str) -> list[str]:
    s = unicodedata.normalize("NFD", texto.lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    parts = re.findall(r"[a-z0-9]+", s)
    return [p for p in parts if p not in STOP and len(p) > 2]


def coincide(nombre: str, encontrado: str) -> bool:
    a, b = set(tokens(nombre)), set(tokens(encontrado))
    if not a or not b:
        return False
    inter = a & b
    if len(inter) >= 2:
        return True
    # marca corta (takis, nescafe, etc.)
    return bool(inter) and (tokens(nombre)[0] in b)


def buscar_nombre(nombre: str) -> str:
    q = " ".join(tokens(nombre)[:5])
    if len(q) < 4:
        return ""
    path = urllib.parse.quote(q)
    for store in STORES[:4]:
        data = get_json(f"{store}/api/catalog_system/pub/products/search/{path}")
        if not isinstance(data, list):
            continue
        for prod in data[:8]:
            titulo = f"{prod.get('productName','')} {prod.get('brand','')}"
            if coincide(nombre, titulo):
                url = imagen_vtex([prod])
                if url:
                    return url
    return ""


def download(url: str, dest: Path) -> bool:
    dest.parent.mkdir(parents=True, exist_ok=True)
    req = urllib.request.Request(
        url,
        headers={"User-Agent": UA, "Accept": "image/avif,image/webp,image/*,*/*"},
    )
    try:
        with urllib.request.urlopen(req, timeout=20) as res:
            data = res.read()
        if len(data) < 1200:
            return False
        dest.write_bytes(data)
        return True
    except Exception:
        return False


def leer_excel():
    wb = openpyxl.load_workbook(EXCEL, data_only=True)
    ws = wb.active
    by_id = {}
    for i, row in enumerate(ws.iter_rows(min_row=11, max_row=ws.max_row, values_only=True), 11):
        nombre = row[3]
        precio = row[8]
        if not nombre or precio is None:
            continue
        barras = str(row[1]).strip() if row[1] else ""
        if barras.lower() == "none":
            barras = ""
        by_id[f"p{i}"] = barras
    return by_id


def main() -> None:
    DEST.mkdir(parents=True, exist_ok=True)
    PUBLIC.mkdir(parents=True, exist_ok=True)
    productos = json.loads(PRODUCTOS.read_text(encoding="utf-8"))
    eans = leer_excel()
    EAN_MAP.write_text(json.dumps(eans, ensure_ascii=False, indent=2), encoding="utf-8")
    cache = {}
    if CACHE.exists():
        try:
            cache = json.loads(CACHE.read_text(encoding="utf-8"))
        except Exception:
            cache = {}

    pendientes = [
        p
        for p in productos
        if not str(p.get("imagen", "")).startswith("http")
        and not str(p.get("imagen", "")).startswith("/recursos/imagenes-productos/")
    ]
    print(f"Sin foto propia: {len(pendientes)} / {len(productos)}")

    def resolver(prod: dict) -> tuple[str, str]:
        pid = prod["id"]
        ean = eans.get(pid, "")
        key = f"ean:{ean}" if ean else f"nom:{prod['nombre']}"
        if cache.get(key):
            return pid, cache[key]
        url = ""
        if ean:
            url = buscar_ean(ean)
        if not url:
            url = buscar_nombre(prod["nombre"])
        cache[key] = url
        return pid, url

    hits = 0
    with ThreadPoolExecutor(max_workers=6) as pool:
        futs = [pool.submit(resolver, p) for p in pendientes]
        total = len(futs)
        for i, fut in enumerate(as_completed(futs), 1):
            pid, url = fut.result()
            if url:
                hits += 1
            if i % 25 == 0 or i == total:
                CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
                print(f"  {i}/{total} urls {hits}")

    CACHE.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")

    bajadas = 0
    for p in productos:
        pid = p["id"]
        ean = eans.get(pid, "")
        key = f"ean:{ean}" if ean else f"nom:{p['nombre']}"
        url = cache.get(key) or ""
        if not url and str(p.get("imagen", "")).startswith("http"):
            url = p["imagen"]
        if not url:
            continue
        dest = DEST / f"{pid}.jpg"
        public = PUBLIC / f"{pid}.jpg"
        if dest.exists() and dest.stat().st_size > 1200:
            ok = True
        else:
            ok = download(url, dest)
        if ok:
            public.write_bytes(dest.read_bytes())
            p["imagen"] = f"/recursos/imagenes-productos/{pid}.jpg"
            bajadas += 1
        else:
            p["imagen"] = url

    PRODUCTOS.write_text(json.dumps(productos, ensure_ascii=False, indent=2), encoding="utf-8")
    propias = sum(
        1
        for p in productos
        if str(p["imagen"]).startswith("/recursos/imagenes-productos/")
        or str(p["imagen"]).startswith("http")
    )
    print(f"Listo. Fotos propias/url: {propias}/{len(productos)} (bajadas {bajadas})")


if __name__ == "__main__":
    main()
